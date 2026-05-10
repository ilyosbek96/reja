from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)

DARK_BG    = RGBColor(0x1E, 0x1E, 0x2E)
ACCENT     = RGBColor(0x89, 0xB4, 0xFA)  # blue
GREEN      = RGBColor(0xA6, 0xE3, 0xA1)
YELLOW     = RGBColor(0xF9, 0xE2, 0xAF)
RED        = RGBColor(0xF3, 0x8B, 0xA8)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
GRAY       = RGBColor(0xBA, 0xC2, 0xDE)
CODE_BG    = RGBColor(0x31, 0x32, 0x44)

def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_textbox(slide, text, left, top, width, height,
                font_size=18, bold=False, color=WHITE,
                align=PP_ALIGN.LEFT, font_name="Consolas"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return txBox

def add_code_box(slide, code, left, top, width, height, font_size=13):
    shape = slide.shapes.add_shape(
        1, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.color.rgb = ACCENT
    shape.line.width = Pt(1.5)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left  = Inches(0.15)
    tf.margin_right = Inches(0.1)
    tf.margin_top   = Inches(0.1)

    first = True
    for line in code.strip().split("\n"):
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_size)
        run.font.name = "Consolas"
        run.font.color.rgb = GREEN

def add_bullet_box(slide, items, left, top, width, height,
                   font_size=16, title=None, title_color=ACCENT):
    shape = slide.shapes.add_shape(
        1, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0x2A, 0x2A, 0x3E)
    shape.line.color.rgb = ACCENT
    shape.line.width = Pt(1)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_top  = Inches(0.12)

    first_p = True
    if title:
        p = tf.paragraphs[0]
        run = p.add_run()
        run.text = title
        run.font.size = Pt(font_size + 1)
        run.font.bold = True
        run.font.color.rgb = title_color
        run.font.name = "Calibri"
        first_p = False

    for item in items:
        if first_p:
            p = tf.paragraphs[0]
            first_p = False
        else:
            p = tf.add_paragraph()
        run = p.add_run()
        run.text = item
        run.font.size = Pt(font_size)
        run.font.color.rgb = WHITE
        run.font.name = "Calibri"


# ─────────────────────────────────────────────
# Slide 1 — Title
# ─────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, DARK_BG)

# decorative accent bar
bar = slide.shapes.add_shape(1, Inches(0), Inches(3.2), Inches(13.33), Inches(0.08))
bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()

add_textbox(slide, "server.js", 0.5, 0.8, 12, 1.2,
            font_size=54, bold=True, color=ACCENT, align=PP_ALIGN.CENTER, font_name="Calibri")
add_textbox(slide, "Node.js + MongoDB Server — Tushuntirish", 0.5, 2.1, 12, 0.7,
            font_size=22, color=GRAY, align=PP_ALIGN.CENTER, font_name="Calibri")
add_textbox(slide, "Umumiy arxitektura va kod tahlili", 0.5, 3.5, 12, 0.6,
            font_size=18, color=WHITE, align=PP_ALIGN.CENTER, font_name="Calibri")

# tech badges
for i, (label, col) in enumerate([("Node.js", GREEN), ("MongoDB", YELLOW), ("Express", ACCENT)]):
    lft = 3.8 + i * 2.2
    badge = slide.shapes.add_shape(1, Inches(lft), Inches(4.5), Inches(1.7), Inches(0.55))
    badge.fill.solid(); badge.fill.fore_color.rgb = CODE_BG
    badge.line.color.rgb = col; badge.line.width = Pt(1.5)
    tb = badge.text_frame
    p = tb.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.size = Pt(15); r.font.bold = True; r.font.color.rgb = col; r.font.name = "Calibri"

add_textbox(slide, "📄  server.js  |  ilyosbek96", 0.5, 6.8, 12, 0.4,
            font_size=12, color=GRAY, align=PP_ALIGN.CENTER, font_name="Calibri")


# ─────────────────────────────────────────────
# Slide 2 — Umumiy oqim (flow)
# ─────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, DARK_BG)
add_textbox(slide, "Dastur ishlash oqimi", 0.4, 0.2, 12, 0.7,
            font_size=30, bold=True, color=ACCENT, font_name="Calibri")

steps = [
    ("1", "Dastur ishga tushadi", GREEN),
    ("2", "MongoDB ga\nulanish", YELLOW),
    ("3", "Muvaffaqiyatli?\nHa / Yo'q", ACCENT),
    ("4", "app.js yuklanadi\nServer 3000-port", GREEN),
]
for i, (num, label, col) in enumerate(steps):
    lx = 0.5 + i * 3.1
    box = slide.shapes.add_shape(1, Inches(lx), Inches(1.4), Inches(2.7), Inches(1.6))
    box.fill.solid(); box.fill.fore_color.rgb = CODE_BG
    box.line.color.rgb = col; box.line.width = Pt(2)
    tf = box.text_frame; tf.margin_left = Inches(0.1); tf.margin_top = Inches(0.1)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = num
    r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = col; r.font.name = "Calibri"
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = label
    r2.font.size = Pt(14); r2.font.color.rgb = WHITE; r2.font.name = "Calibri"

    if i < 3:
        arr = slide.shapes.add_shape(1, Inches(lx + 2.7), Inches(1.9), Inches(0.4), Inches(0.5))
        arr.fill.solid(); arr.fill.fore_color.rgb = ACCENT; arr.line.fill.background()

add_code_box(slide,
"""Dastur ishga tushdi
    → mongodb.connect(connectionString)
        ✅ Muvaffaqiyatli:
            → module.exports = client
            → require("./app")
            → server.listen(3000)
        ❌ Xato:
            → "MongoDB ULANMADI" xabari
            → Server ishlamaydi""",
0.5, 3.2, 12.3, 3.8, font_size=14)


# ─────────────────────────────────────────────
# Slide 3 — Modullar
# ─────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, DARK_BG)
add_textbox(slide, "1. Modullarni import qilish", 0.4, 0.2, 12, 0.6,
            font_size=28, bold=True, color=ACCENT, font_name="Calibri")

add_code_box(slide,
"""const http    = require("http");
const mongodb = require("mongodb");""",
0.4, 1.0, 12.5, 1.2)

add_bullet_box(slide,
    ["🔵  http  —  Node.js ning ichki moduli (o'rnatish shart emas)",
     "           HTTP server yaratish uchun ishlatiladi",
     "",
     "🟢  mongodb  —  tashqi kutubxona (npm install mongodb)",
     "           MongoDB baza bilan bog'lanish uchun ishlatiladi"],
    0.4, 2.5, 12.5, 2.5, font_size=16)

add_textbox(slide,
    "💡  Node.js da har bir fayl o'z moduliga ega — require() orqali ularga murojaat qilinadi.",
    0.4, 5.2, 12.5, 0.9, font_size=15, color=YELLOW, font_name="Calibri")


# ─────────────────────────────────────────────
# Slide 4 — Connection String
# ─────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, DARK_BG)
add_textbox(slide, "2. MongoDB ulanish manzili (Connection String)", 0.4, 0.2, 12, 0.6,
            font_size=26, bold=True, color=ACCENT, font_name="Calibri")

add_code_box(slide,
"""const connectionString =
  "mongodb+srv://ilyosbekk72_db_user:roLWHdiuusH5nTdf
   @cluster0.lnmyuud.mongodb.net/Reja?appName=Cluster0";""",
0.4, 1.0, 12.5, 1.5)

parts = [
    ("mongodb+srv://", "Protokol\n(bulut ulanish)", ACCENT),
    ("ilyosbekk72_db_user", "Foydalanuvchi\nnomi", GREEN),
    ("roLWHdiuusH5nTdf", "Parol\n(maxfiy!)", RED),
    ("cluster0.lnmyuud\n.mongodb.net", "Cluster\nmanzili", YELLOW),
    ("/Reja", "Baza\nnomi", ACCENT),
]
for i, (val, desc, col) in enumerate(parts):
    lx = 0.3 + i * 2.5
    b = slide.shapes.add_shape(1, Inches(lx), Inches(2.8), Inches(2.3), Inches(1.5))
    b.fill.solid(); b.fill.fore_color.rgb = CODE_BG
    b.line.color.rgb = col; b.line.width = Pt(1.5)
    tf = b.text_frame; tf.margin_left = Inches(0.08); tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = val
    r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = col; r.font.name = "Consolas"
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = desc
    r2.font.size = Pt(12); r2.font.color.rgb = GRAY; r2.font.name = "Calibri"

# warning box
warn = slide.shapes.add_shape(1, Inches(0.4), Inches(4.6), Inches(12.5), Inches(1.3))
warn.fill.solid(); warn.fill.fore_color.rgb = RGBColor(0x45, 0x26, 0x26)
warn.line.color.rgb = RED; warn.line.width = Pt(2)
tf = warn.text_frame; tf.margin_left = Inches(0.2); tf.margin_top = Inches(0.15)
p = tf.paragraphs[0]
r = p.add_run(); r.text = "⚠️  XAVFSIZLIK MUAMMOSI"
r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = RED; r.font.name = "Calibri"
p2 = tf.add_paragraph()
r2 = p2.add_run()
r2.text = "Parol kodda ochiq ko'rinib turibdi! Uni .env faylga ko'chirish kerak:  DB_PASS=roLWHdiuusH5nTdf"
r2.font.size = Pt(14); r2.font.color.rgb = YELLOW; r2.font.name = "Calibri"


# ─────────────────────────────────────────────
# Slide 5 — mongodb.connect()
# ─────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, DARK_BG)
add_textbox(slide, "3. MongoDB ga ulanish", 0.4, 0.2, 12, 0.6,
            font_size=28, bold=True, color=ACCENT, font_name="Calibri")

add_code_box(slide,
"""mongodb.connect(
  connectionString,
  { useNewUrlParser: true, useUnifiedTopology: true },
  (err, client) => {
    if (err) console.log("ERROR  MongoDB ULANMADI");
    else {
      console.log("MongoDB connection succeed");
      module.exports = client;
      ...
    }
  }
);""",
0.4, 1.0, 6.0, 4.5)

add_bullet_box(slide, [
    "📌  3 ta argument qabul qiladi:",
    "",
    "  1️⃣  connectionString — baza manzili",
    "  2️⃣  options — sozlamalar (eski ogohlantirishlarni o'chirish)",
    "  3️⃣  callback — natija (err yoki client)",
    "",
    "✅  err = null  →  ulanish muvaffaqiyatli",
    "❌  err = Error →  ulanish amalga oshmadi",
    "",
    "module.exports = client  →  boshqa fayllarda",
    "  bazadan foydalanish uchun export qilinadi",
], 6.7, 1.0, 6.2, 4.5, font_size=14)


# ─────────────────────────────────────────────
# Slide 6 — app.js va Server
# ─────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, DARK_BG)
add_textbox(slide, "4. app.js yuklanishi va Server ishga tushishi", 0.4, 0.2, 12.5, 0.6,
            font_size=26, bold=True, color=ACCENT, font_name="Calibri")

add_code_box(slide,
"""module.exports = client;         // ← client ni export qil

const app    = require("./app"); // ← app.js ni yukla
const server = http.createServer(app);

let PORT = 3000;
server.listen(PORT, function () {
  console.log(`Server ishlayapti: http://localhost:${PORT}`);
});""",
0.4, 1.0, 12.5, 3.2)

add_bullet_box(slide, [
    "🔑  Nima uchun app.js KEYIN yuklanadi?",
    "   → app.js bazaga murojaat qilishi mumkin.",
    "   → Agar avval yuklanadigan bo'lsa, baza tayyor bo'lmay xato beradi.",
    "",
    "🌐  http.createServer(app)  —  Express app ni HTTP serverga o'rash",
    "🔌  server.listen(3000)     —  3000-portda so'rovlarni kutib olish boshlash",
], 0.4, 4.4, 12.5, 2.7, font_size=15)


# ─────────────────────────────────────────────
# Slide 7 — Xulosa
# ─────────────────────────────────────────────
slide = prs.slides.add_slide(prs.slide_layouts[6])
set_bg(slide, DARK_BG)
add_textbox(slide, "Xulosa", 0.4, 0.2, 12, 0.6,
            font_size=34, bold=True, color=ACCENT, font_name="Calibri")

items = [
    ("✅", "http & mongodb modullar import qilinadi",          GREEN),
    ("✅", "MongoDB Atlas ga ulanish amalga oshiriladi",        GREEN),
    ("✅", "Muvaffaqiyatli bo'lsa — app.js yuklanadi",          GREEN),
    ("✅", "HTTP server yaratiladi va 3000-portda ishga tushadi",GREEN),
    ("⚠️", "Parol kodda ochiq — .env faylga ko'chirish kerak", YELLOW),
    ("💡", "app.js faqat baza ulangach yuklanadi — to'g'ri tartib",ACCENT),
]
for i, (icon, text, col) in enumerate(items):
    row = slide.shapes.add_shape(1, Inches(0.4), Inches(1.1 + i*0.9), Inches(12.5), Inches(0.75))
    row.fill.solid(); row.fill.fore_color.rgb = CODE_BG
    row.line.color.rgb = col; row.line.width = Pt(1)
    tf = row.text_frame; tf.margin_left = Inches(0.15); tf.margin_top = Inches(0.08)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = f"{icon}  {text}"
    r.font.size = Pt(16); r.font.color.rgb = col; r.font.name = "Calibri"

add_textbox(slide, "server.js  |  Node.js + MongoDB  |  ilyosbek96",
            0.4, 7.0, 12.5, 0.35, font_size=12, color=GRAY,
            align=PP_ALIGN.CENTER, font_name="Calibri")


# ─────────────────────────────────────────────
prs.save("/Users/ilyosbek96/Desktop/serverjs_tushuntirish.pptx")
print("✅ PPT yaratildi: Desktop/serverjs_tushuntirish.pptx")
